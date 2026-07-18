"""
UHRP document ledger — an Excel workbook recording every file uploaded to
or downloaded from UHRP storage: filename, hash, the uhrp:// URL, a
clickable direct link to fetch the file, and an embedded thumbnail for
image files.

Stored at ~/.bestbrand/uhrp_ledger.xlsx, alongside registry.json (registry.py)
and fees.jsonl — this app's existing local-data convention.
"""
from __future__ import annotations
import io
import time
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from loguru import logger

_BBC_DIR = Path.home() / ".bestbrand"
LEDGER_PATH = _BBC_DIR / "uhrp_ledger.xlsx"
FILES_DIR = _BBC_DIR / "uhrp_files"

HEADERS = [
    "Date",
    "Direction",
    "Filename",
    "Content Type",
    "Size (bytes)",
    "SHA-256 Hash",
    "UHRP URL",
    "Direct Link",
    "Expiry Time",
    "Status",
    "Hosted By",
    "Local Path",
    "Requested By",
    "Thumbnail",
]
COL = {name: i + 1 for i, name in enumerate(HEADERS)}
THUMB_COL_LETTER = get_column_letter(COL["Thumbnail"])
THUMB_SIZE = 64  # px, square

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}


def ensure_dirs():
    _BBC_DIR.mkdir(parents=True, exist_ok=True)
    FILES_DIR.mkdir(parents=True, exist_ok=True)


def unique_local_path(filename: str) -> Path:
    """
    Returns a path under FILES_DIR for filename, disambiguating with a
    numeric suffix if a file by that name already exists — otherwise a
    later upload/download reusing a common filename (e.g. "image.jpg")
    would silently overwrite an earlier one still referenced by the ledger.
    """
    ensure_dirs()
    candidate = FILES_DIR / filename
    if not candidate.exists():
        return candidate
    stem, suffix = Path(filename).stem, Path(filename).suffix
    n = 1
    while True:
        candidate = FILES_DIR / f"{stem}_{n}{suffix}"
        if not candidate.exists():
            return candidate
        n += 1


def _new_workbook() -> Workbook:
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "UHRP Ledger"
    ws.append(HEADERS)
    for i, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = cell.font.copy(bold=True)
    widths = [17, 10, 28, 22, 12, 20, 30, 30, 17, 10, 24, 34, 16, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return wb


def _open_workbook() -> Workbook:
    ensure_dirs()
    if LEDGER_PATH.exists():
        return load_workbook(LEDGER_PATH)
    wb = _new_workbook()
    wb.save(LEDGER_PATH)
    return wb


def get_ledger_path() -> Path:
    _open_workbook()
    return LEDGER_PATH


def _thumbnail_bytes(local_path: str) -> Optional[bytes]:
    ext = Path(local_path).suffix.lower()
    if ext not in IMAGE_EXTS:
        return None
    try:
        from PIL import Image
        img = Image.open(local_path)
        img.thumbnail((THUMB_SIZE, THUMB_SIZE))
        buf = io.BytesIO()
        img.convert("RGB").save(buf, format="PNG")
        return buf.getvalue()
    except Exception as e:
        logger.warning(f"UHRP ledger: could not build thumbnail for {local_path}: {e}")
        return None


def _format_expiry(expiry_time: Optional[int]) -> str:
    if not expiry_time:
        return ""
    return time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(expiry_time))


def _status_for(expiry_time: Optional[int]) -> str:
    if not expiry_time:
        return ""
    return "Active" if expiry_time > time.time() else "Expired"


def add_entry(
    direction: str,
    filename: str,
    content_type: str,
    size: int,
    sha256: str,
    uhrp_url: str,
    direct_link: Optional[str],
    hosted_by: list[str],
    local_path: Optional[str],
    requested_by: str,
    expiry_time: Optional[int] = None,
) -> int:
    """Appends one row to the ledger. Returns the row number written."""
    wb = _open_workbook()
    ws: Worksheet = wb["UHRP Ledger"]
    row = ws.max_row + 1

    values = [
        time.strftime("%Y-%m-%d %H:%M:%S", time.localtime()),
        direction,
        filename,
        content_type,
        size,
        sha256,
        uhrp_url,
        direct_link or "(pending — resolve to get link)",
        _format_expiry(expiry_time),
        _status_for(expiry_time),
        ", ".join(hosted_by) if hosted_by else "",
        local_path or "",
        requested_by,
    ]
    for col_idx, val in enumerate(values, start=1):
        ws.cell(row=row, column=col_idx, value=val)

    link_cell = ws.cell(row=row, column=COL["Direct Link"])
    if direct_link:
        link_cell.hyperlink = direct_link
        link_cell.style = "Hyperlink"

    if local_path:
        thumb = _thumbnail_bytes(local_path)
        if thumb:
            img = XLImage(io.BytesIO(thumb))
            img.width = THUMB_SIZE
            img.height = THUMB_SIZE
            ws.row_dimensions[row].height = THUMB_SIZE * 0.75
            ws.add_image(img, f"{THUMB_COL_LETTER}{row}")

    wb.save(LEDGER_PATH)
    return row


def update_direct_link(uhrp_url: str, direct_link: str) -> int:
    """Backfills the Direct Link column for every row matching uhrp_url. Returns rows updated."""
    wb = _open_workbook()
    ws: Worksheet = wb["UHRP Ledger"]
    updated = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=COL["UHRP URL"]).value == uhrp_url:
            cell = ws.cell(row=row, column=COL["Direct Link"], value=direct_link)
            cell.hyperlink = direct_link
            cell.style = "Hyperlink"
            updated += 1
    if updated:
        wb.save(LEDGER_PATH)
    return updated


def known_uhrp_urls() -> set[str]:
    """All UHRP URLs already present in the ledger — used to skip duplicates when syncing."""
    wb = _open_workbook()
    ws: Worksheet = wb["UHRP Ledger"]
    return {
        ws.cell(row=r, column=COL["UHRP URL"]).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=COL["UHRP URL"]).value
    }


def update_expiry(uhrp_url: str, expiry_time: int) -> int:
    """Backfills Expiry Time + Status for every row matching uhrp_url. Returns rows updated."""
    wb = _open_workbook()
    ws: Worksheet = wb["UHRP Ledger"]
    updated = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=COL["UHRP URL"]).value == uhrp_url:
            ws.cell(row=row, column=COL["Expiry Time"], value=_format_expiry(expiry_time))
            ws.cell(row=row, column=COL["Status"], value=_status_for(expiry_time))
            updated += 1
    if updated:
        wb.save(LEDGER_PATH)
    return updated


def recent_entries(limit: int = 10) -> list[dict]:
    wb = _open_workbook()
    ws: Worksheet = wb["UHRP Ledger"]
    rows = []
    for row in range(2, ws.max_row + 1):
        rows.append({h: ws.cell(row=row, column=COL[h]).value for h in HEADERS if h != "Thumbnail"})
    return rows[-limit:][::-1]
